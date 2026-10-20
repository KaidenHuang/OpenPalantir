"""② 表格提取单元测试 —— 表档案确定性渲染、边界与失败模式（纯解析，无 LLM、无外部服务）

覆盖 docs/data-flow.md「表格文件分支」② 一节的判据：
1. 注册表含 .xlsx .xlsm .csv .tsv，.xls 给可读错误
2. 同一文件两次 extract_content 逐字节相同
3. 每条边界各有用例并留痕
4. iter_table_rows 的列名集合与表档案逐名一致（③ 的标注能在 ④ 对上）
"""
import datetime
import json
import os
import sys

import pytest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND = os.path.join(BASE, "backend")
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from content_extraction.document import extract_content, iter_table_rows      # noqa: E402
from content_extraction.document.converters import supported_extensions       # noqa: E402
from content_extraction.document.models import KIND_TABLE                     # noqa: E402
from content_extraction.document.table_model import (                         # noqa: E402
    MAX_PROFILE_COLUMNS, MAX_SHEETS, SheetTable, TYPE_BOOL, TYPE_DATE, TYPE_EMPTY,
    TYPE_NUM, TYPE_TEXT, infer_column_type, render_cell, render_profile,
)
from content_extraction.document.table_reader import CSV_SHEET_TITLE          # noqa: E402

pytestmark = pytest.mark.unit


def write_csv(path, text, encoding="utf-8"):
    path.write_bytes(text.encode(encoding))
    return str(path)


def workbook(path, sheets):
    """sheets: {标题: 行列表}，行元素按 openpyxl 原样写入"""
    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    for index, (title, rows) in enumerate(sheets.items()):
        ws = first if index == 0 else wb.create_sheet()
        ws.title = title
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return str(path)


class TestRegistry:
    def test_table_extensions_registered(self):
        for ext in (".xlsx", ".xlsm", ".csv", ".tsv"):
            assert ext in supported_extensions()

    def test_xls_rejected_with_hint(self, tmp_path):
        legacy = tmp_path / "老表.xls"
        legacy.write_bytes(b"not a real xls")
        with pytest.raises(ValueError) as exc:
            extract_content(str(legacy))
        assert "另存为" in str(exc.value)


class TestXlsx:
    def test_one_sheet_one_level1_section(self, tmp_path):
        path = workbook(tmp_path / "销售统计.xlsx", {
            "月度汇总": [["部门", "省份", "销售额", "月份"],
                       ["华东大区", "浙江", 1203400.5, datetime.date(2025, 1, 1)],
                       ["华北大区", "北京", 980200, datetime.date(2025, 1, 1)]],
            "人员": [["工号", "姓名"], [10001, "张伟"]],
        })
        doc = extract_content(path)
        assert doc.format == "xlsx"
        assert [s.title for s in doc.sections] == ["月度汇总", "人员"]
        assert all(s.level == 1 and s.kind == KIND_TABLE for s in doc.sections)
        assert all(s.page_start is None for s in doc.sections)
        text = doc.sections[0].text
        assert text.startswith("【工作表】月度汇总")
        assert "共 2 行 × 4 列" in text
        assert "部门(TEXT)" in text and "销售额(NUM)" in text and "月份(DATE)" in text

    def test_sample_rows_are_first_two(self, tmp_path):
        rows = [["编号", "名称"]] + [[i, f"第{i}行"] for i in range(1, 6)]
        doc = extract_content(workbook(tmp_path / "many.xlsx", {"表": rows}))
        text = doc.sections[0].text
        assert "第1行" in text and "第2行" in text and "第3行" not in text
        assert "共 5 行" in text
        assert "仅展示前 2 行，共 5 行" in text

    def test_hidden_sheet_marked(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "内部"
        ws.append(["列A", "列B"])
        ws.append([1, 2])
        ws.sheet_state = "hidden"
        # openpyxl 不允许唯一工作表是隐藏的（Excel 本身的规则），补一个可见表
        other = wb.create_sheet("可见")
        other.append(["列A", "列B"])
        other.append([3, 4])
        path = str(tmp_path / "hidden.xlsx")
        wb.save(path)
        doc = extract_content(path)
        assert "(隐藏工作表)" in doc.sections[0].text

    def test_merged_cells_counted_and_filled_once(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "合并"
        ws.merge_cells("A1:B1")
        ws["A1"] = "合计标题"
        ws["A3"] = "部门"
        ws["B3"] = "金额"
        ws["A4"] = "销售"
        ws["B4"] = 100
        path = str(tmp_path / "merged.xlsx")
        wb.save(path)
        doc = extract_content(path)
        assert "合并单元格 1 处" in doc.sections[0].text

    def test_formula_cells_marked(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "公式"
        ws.append(["数量", "单价", "金额"])
        ws.append([3, 10, "=A2*B2"])
        ws.append([5, 7, "=A3*B3"])
        path = str(tmp_path / "formula.xlsx")
        wb.save(path)
        doc = extract_content(path)
        text = doc.sections[0].text
        assert "公式单元格 2 处" in text
        assert "#FORMULA" in text
        # ④ 的行数据里公式不是取值，是空串
        rows = list(iter_table_rows(path, "公式"))
        assert rows[0]["金额"] == ""

    def test_empty_sheet_skipped_with_reason(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "有数据"
        wb.active.append(["列A", "列B"])
        wb.active.append([1, 2])
        wb.create_sheet("空表")
        path = str(tmp_path / "empty_sheet.xlsx")
        wb.save(path)
        doc = extract_content(path)
        assert [s.title for s in doc.sections] == ["有数据"]
        assert {"name": "空表", "reason": "empty"} in doc.meta["skipped_sheets"]

    def test_single_column_sheet_skipped(self, tmp_path):
        path = workbook(tmp_path / "one_col.xlsx", {
            "正常": [["甲", "乙"], [1, 2]],
            "只有一列": [["项目"], ["甲"], ["乙"]],
        })
        doc = extract_content(path)
        assert [s.title for s in doc.sections] == ["正常"]
        assert {"name": "只有一列", "reason": "single_column"} in doc.meta["skipped_sheets"]

    def test_all_sheets_unusable_raises(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "空"
        path = str(tmp_path / "全空.xlsx")
        wb.save(path)
        with pytest.raises(ValueError):
            extract_content(path)

    def test_more_than_max_sheets(self, tmp_path):
        sheets = {f"表{i}": [["A", "B"], [i, i * 2]] for i in range(35)}
        doc = extract_content(workbook(tmp_path / "many_sheets.xlsx", sheets))
        assert len(doc.sections) == MAX_SHEETS
        limited = [item for item in doc.meta["skipped_sheets"] if item["reason"] == "sheet_limit"]
        assert len(limited) == 5

    def test_wide_sheet_profiles_first_columns(self, tmp_path):
        columns = [f"列{i}" for i in range(1, 201)]
        doc = extract_content(workbook(tmp_path / "wide.xlsx", {"宽表": [columns, [1] * 200]}))
        text = doc.sections[0].text
        assert "另有 170 列" in text
        assert text.count("(TEXT)") + text.count("(NUM)") <= MAX_PROFILE_COLUMNS

    def test_byte_stable_render(self, tmp_path):
        path = workbook(tmp_path / "stable.xlsx", {
            "表": [["部门", "金额", "日期"],
                  ["华东", 0.1 + 0.2, datetime.datetime(2025, 1, 1, 8, 30)],
                  ["华北", 3, datetime.datetime(2025, 2, 1)]]})

        def dump(doc):
            return json.dumps([section.__dict__ for section in doc.sections],
                              sort_keys=True, ensure_ascii=False, default=str)

        assert dump(extract_content(path)) == dump(extract_content(path))


class TestCsv:
    def test_gbk_without_bom(self, tmp_path):
        path = write_csv(tmp_path / "花名册.csv",
                         "工号,姓名,部门\n10001,张伟,销售\n", encoding="gbk")
        doc = extract_content(path)
        assert "姓名" in doc.sections[0].text
        assert "张伟" in doc.sections[0].text
        assert doc.format == "csv"
        assert doc.sections[0].title == CSV_SHEET_TITLE

    def test_semicolon_and_pipe_delimiters(self, tmp_path):
        for name, delimiter in (("sep_semi", ";"), ("sep_pipe", "|")):
            path = write_csv(tmp_path / f"{name}.csv",
                             f"大区{delimiter}省份\n华东{delimiter}浙江\n")
            text = extract_content(path).sections[0].text
            assert "大区(TEXT)" in text and "省份(TEXT)" in text
            assert "共 1 行 × 2 列" in text

    def test_tsv_format(self, tmp_path):
        path = write_csv(tmp_path / "表格.tsv", "甲\t乙\n1\t2\n")
        doc = extract_content(path)
        assert doc.format == "tsv"
        assert "共 1 行 × 2 列" in doc.sections[0].text

    def test_ragged_rows_padded_and_counted(self, tmp_path):
        path = write_csv(tmp_path / "ragged.csv",
                         "工号,姓名,部门\n1,张伟\n2,李四,销售,备注追加\n")
        text = extract_content(path).sections[0].text
        assert "非法列数行 2 条" in text
        assert "备注追加" in text

    def test_header_detected_below_title_row(self, tmp_path):
        path = write_csv(tmp_path / "带标题.csv",
                         "2025年销售统计,,,,\n部门,省份,销售额\n华东,浙江,100\n")
        doc = extract_content(path)
        text = doc.sections[0].text
        assert "表头取自第 2 行" in text
        assert "部门(TEXT)" in text and "2025年销售统计" not in text

    def test_header_failure_uses_placeholders(self, tmp_path):
        # 每行只有首格有值（非空率 1/4 不合格），前 5 行都当不了表头 → 占位列名
        path = write_csv(tmp_path / "无表头.csv", "销售额,,,\n成本,,,\n毛利,,,\n")
        text = extract_content(path).sections[0].text
        assert "列1(TEXT)" in text and "列2(EMPTY)" in text
        assert "表头识别失败（用占位列名）" in text

    def test_numeric_header_kept_as_names(self, tmp_path):
        # 年份当列名的宽表是真实形态：识别得到的行就照它命名，不占位
        path = write_csv(tmp_path / "年份列.csv", "指标,2023,2024\n收入,100,120\n")
        text = extract_content(path).sections[0].text
        assert "2023(NUM)" in text and "表头识别失败" not in text

    def test_duplicate_column_names_stay_addressable(self, tmp_path):
        path = write_csv(tmp_path / "重名.csv", "金额,金额,备注\n1,2,x\n")
        text = extract_content(path).sections[0].text
        assert "金额(NUM)" in text and "金额_2(NUM)" in text
        assert list(list(iter_table_rows(path, CSV_SHEET_TITLE))[0].keys()) == ["金额", "金额_2", "备注"]

    def test_single_column_file_rejected(self, tmp_path):
        path = write_csv(tmp_path / "其实是txt.csv", "第一句话\n第二句话\n")
        with pytest.raises(ValueError) as exc:
            extract_content(path)
        assert "single_column" in str(exc.value)

    @pytest.mark.slow
    def test_rows_capped_and_marked(self, tmp_path):
        lines = ["编号,名称"] + [f"{i},第{i}行" for i in range(1, 50002)]
        path = write_csv(tmp_path / "big.csv", "\n".join(lines) + "\n")
        text = extract_content(path).sections[0].text
        assert "已按前 50000 行统计" in text
        assert "共 50000 行" in text

    def test_iter_table_rows_matches_profile_columns(self, tmp_path):
        path = write_csv(tmp_path / "对齐.csv", "部门,月份,销售额\n华东,2025-01,100\n华北,2025-02,200\n")
        rows = list(iter_table_rows(path, CSV_SHEET_TITLE))
        assert [r["部门"] for r in rows] == ["华东", "华北"]
        assert rows[1] == {"部门": "华北", "月份": "2025-02", "销售额": "200"}

    def test_unknown_sheet_title_yields_nothing(self, tmp_path):
        path = workbook(tmp_path / "两个表.xlsx", {"甲表": [["A", "B"], [1, 2]],
                                                 "乙表": [["C", "D"], [3, 4]]})
        assert list(iter_table_rows(path, "不存在的表")) == []
        assert list(iter_table_rows(path, "乙表")) == [{"C": "3", "D": "4"}]


class TestTableModel:
    @pytest.mark.parametrize("value,expected", [
        (None, ""),
        (True, "TRUE"),
        (False, "FALSE"),
        (3, "3"),
        (0.1 + 0.2, "0.3"),
        (1234567.0, "1.23457e+06"),
        (datetime.date(2025, 1, 5), "2025-01-05"),
        (datetime.datetime(2025, 1, 5, 8, 30, 9), "2025-01-05 08:30:09"),
        ("  带换行\n文本  ", "带换行 文本"),
        ("含|竖线", "含／竖线"),
    ])
    def test_render_cell_rules(self, value, expected):
        assert render_cell(value) == expected

    def test_long_cell_truncated(self):
        assert render_cell("甲" * 300).endswith("…")
        assert len(render_cell("甲" * 300)) == 201

    @pytest.mark.parametrize("values,expected", [
        ([1, 2, 3], TYPE_NUM),
        (["1", "2.5"], TYPE_NUM),
        ([datetime.date(2025, 1, 1)], TYPE_DATE),
        ([True, False], TYPE_BOOL),
        (["张伟", "李四"], TYPE_TEXT),
        ([None, ""], TYPE_EMPTY),
        ([1, "张伟"], TYPE_TEXT),
    ])
    def test_infer_column_type(self, values, expected):
        assert infer_column_type(values) == expected

    def test_profile_omits_sample_block_without_rows(self):
        table = SheetTable(title="空表", columns=["甲", "乙"], column_types=[TYPE_TEXT, TYPE_NUM],
                           row_count=0)
        assert "【样本】" not in render_profile(table)

    def test_preview_rows_can_be_trimmed(self):
        table = SheetTable(title="表", columns=["甲", "乙"], column_types=[TYPE_TEXT, TYPE_NUM],
                           row_count=9, samples=[["a", "1"], ["b", "2"]])
        assert "b" in render_profile(table, preview_rows=2)
        assert "b" not in render_profile(table, preview_rows=1)
        assert "【样本】" not in render_profile(table, preview_rows=0)
