"""content_extraction 单元测试 —— 转换器与章节树解析（纯解析，无 LLM、无外部服务）"""
import os
import sys

import pytest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND = os.path.join(BASE, "backend")
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from content_extraction.document import extract_content, ExtractedDocument, Section          # noqa: E402
from content_extraction.document.converters import get_converter, supported_extensions       # noqa: E402
from content_extraction.document.md_parser import parse_markdown_sections, single_section    # noqa: E402


# ── 注册表 ──────────────────────────────────────────────────

def test_supported_extensions():
    from content_extraction.document.converters import supported_extensions
    exts = supported_extensions()
    for ext in (".pdf", ".docx", ".pptx", ".md", ".txt"):
        assert ext in exts
        assert get_converter(ext) is not None

def test_unsupported_extension():
    assert get_converter(".ppt") is None
    with pytest.raises(ValueError):
        extract_content("不存在路径.ppt")


# ── Markdown 章节解析 ───────────────────────────────────────

def test_md_hierarchy_and_text():
    md = "# 主标题\n\n引言正文\n\n## 小节A\n\nA 的正文\n\n### 细节\n\n细节正文\n\n## 小节B\n\nB 的正文\n"
    sections = parse_markdown_sections([(None, md)])
    assert len(sections) == 1
    root = sections[0]
    assert root.title == "主标题"
    assert "引言正文" in root.text
    assert [c.title for c in root.children] == ["小节A", "小节B"]
    a = root.children[0]
    assert a.level == 2
    assert any(c.title == "细节" for c in a.children)
    # 正文归属：A 的正文在 A、细节正文在细节、B 的正文在 B，互不串
    assert "A 的正文" in a.text and "细节正文" not in a.text
    assert "细节正文" in a.children[0].text
    assert "B 的正文" in root.children[1].text


def test_md_preamble_before_first_heading():
    md = "文档开头的导语\n\n# 标题\n\n正文\n"
    sections = parse_markdown_sections([(None, md)])
    assert sections[0].title == "（开头）"
    assert "导语" in sections[0].text
    # 随后的 H1 是同一层级，与「开头」并列
    assert [c.title for c in sections[0].children] == []
    assert len(sections) == 2 and sections[1].title == "标题"


def test_md_code_fence_not_parsed_as_heading():
    md = "## 小节\n\n```python\n# 这不是标题\n```\n\n后续\n"
    sections = parse_markdown_sections([(None, md)])
    assert len(sections) == 1
    assert "# 这不是标题" in sections[0].text


def test_md_page_anchor():
    pages = [(0, "# 首页标题\n\n首页正文\n"), (1, "跨页正文\n"), (2, "# 次页标题\n\n次页正文\n")]
    sections = parse_markdown_sections(pages)
    assert sections[0].page_start == 1 and sections[0].page_end == 2
    assert sections[1].page_start == 3 and sections[1].page_end == 3


def test_single_section_fallback():
    sections = single_section("doc", "全部正文")
    assert sections[0].title == "doc"
    assert "全部正文" in sections[0].text


# ── TXT / DOCX ─────────────────────────────────────────────

def test_txt_converter(tmp_path):
    p = tmp_path / "sample.txt"
    p.write_text("第一段\n\n第二段", encoding="utf-8")
    ex = extract_content(str(p))
    assert ex.format == "txt"
    assert ex.page_count is None
    assert len(ex.sections) == 1
    assert "第二段" in ex.sections[0].text


def _make_docx(tmp_path, body: str):
    from docx import Document

    doc = Document()
    doc.add_heading("文档标题", level=1)
    doc.add_paragraph("导语段落")
    doc.add_heading("方案一", level=2)
    doc.add_paragraph(body)
    doc.add_heading("方案二", level=2)
    doc.add_paragraph("方案二正文")
    p = tmp_path / "sample.docx"
    doc.save(str(p))
    return str(p)


def test_docx_converter(tmp_path):
    path = _make_docx(tmp_path, "方案一正文")
    ex = extract_content(path)
    assert ex.format == "docx"
    assert [s.title for s in ex.sections] == ["文档标题"]
    root = ex.sections[0]
    assert "导语段落" in root.text
    assert [c.title for c in root.children] == ["方案一", "方案二"]
    assert "方案一正文" in root.children[0].text


def test_docx_table_content(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("标题", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "姓名"
    table.rows[0].cells[1].text = "部门"
    p = tmp_path / "with_table.docx"
    doc.save(str(p))
    ex = extract_content(str(p))
    assert "姓名" in ex.sections[0].text and "部门" in ex.sections[0].text


def test_docx_empty_raises(tmp_path):
    from docx import Document

    p = tmp_path / "empty.docx"
    Document().save(str(p))
    with pytest.raises(ValueError):
        extract_content(str(p))


# ── PPTX ───────────────────────────────────────────────────


def _make_pptx(tmp_path, name: str = "sample.pptx") -> str:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()

    content = prs.slides.add_slide(prs.slide_layouts[1])
    content.shapes.title.text = "季度营收回顾"
    body = content.placeholders[1].text_frame
    body.text = "总体增长"
    bullet = body.add_paragraph()
    bullet.text = "华东区同比增长 18%"
    bullet.level = 1

    title_only = prs.slides.add_slide(prs.slide_layouts[5])
    title_only.shapes.title.text = "结语"

    blank = prs.slides.add_slide(prs.slide_layouts[6])
    table = blank.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "部门"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = "研发"

    with_notes = prs.slides.add_slide(prs.slide_layouts[5])
    with_notes.shapes.title.text = "交付计划"
    with_notes.notes_slide.notes_text_frame.text = "李四负责交付。"

    path = tmp_path / name
    prs.save(str(path))
    return str(path)


def test_pptx_converter(tmp_path):
    ex = extract_content(_make_pptx(tmp_path))
    assert ex.format == "pptx"
    assert ex.page_count == 4
    # 无标题占位符的表格页退化为「第 N 页」
    assert [s.title for s in ex.sections] == ["季度营收回顾", "结语", "第 3 页", "交付计划"]
    assert [(s.page_start, s.page_end) for s in ex.sections] == [(1, 1), (2, 2), (3, 3), (4, 4)]

    first = ex.sections[0]
    assert "总体增长" in first.text
    assert "  - 华东区同比增长 18%" in first.text
    # 纯标题页正文回填标题，否则该章节在摘要树与抽取批次中都会消失
    assert ex.sections[1].text == "结语"
    assert "| 姓名 | 部门 |" in ex.sections[2].text
    assert "（备注）李四负责交付。" in ex.sections[3].text


def test_pptx_uppercase_extension(tmp_path):
    """真实文件常见全大写扩展名（.PPTX），注册表按小写查找"""
    ex = extract_content(_make_pptx(tmp_path, name="SAMPLE.PPTX"))
    assert ex.format == "pptx"


def test_pptx_title_from_first_line(tmp_path):
    """企业模板常把标题画成普通文本框：无标题占位符时取第一行文本作章节标题"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for offset, content in enumerate(["数据可用性战略", "围绕数据的有用性管理展开"], start=1):
        box = slide.shapes.add_textbox(Inches(1), Inches(offset), Inches(6), Inches(1))
        box.text_frame.text = content
    path = tmp_path / "template.pptx"
    prs.save(str(path))

    ex = extract_content(str(path))
    assert ex.sections[0].title == "数据可用性战略"
    assert "数据可用性战略" not in ex.sections[0].text
    assert "围绕数据的有用性管理展开" in ex.sections[0].text


def test_pptx_table_only_slide_keeps_page_fallback(tmp_path):
    """表格行不能当标题，仍退化为页码"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "仅表格"
    path = tmp_path / "table_only.pptx"
    prs.save(str(path))

    ex = extract_content(str(path))
    assert ex.sections[0].title == "第 1 页"
    assert "| 仅表格 |" in ex.sections[0].text


def test_pptx_group_shape_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "组合图形内的文本"
    path = tmp_path / "group.pptx"
    prs.save(str(path))

    ex = extract_content(str(path))
    assert "组合图形内的文本" in ex.sections[0].text


def test_pptx_empty_raises(tmp_path):
    from pptx import Presentation

    path = tmp_path / "empty.pptx"
    Presentation().save(str(path))
    with pytest.raises(ValueError):
        extract_content(str(path))


def test_pptx_corrupt_file_raises_value_error(tmp_path):
    """损坏文件与旧版 .ppt 二进制都会被 python-pptx 拒为 PackageNotFoundError，须转成 ValueError"""
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"not an office open xml package")
    with pytest.raises(ValueError):
        extract_content(str(path))


def test_pptx_chart_text(tmp_path):
    """图表文字进正文：类别 × 系列表格，数值原样输出不做单位推断"""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["第一季度", "第二季度"]
    data.add_series("销售额", (0.48, 0.25))
    slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,
                           Inches(1), Inches(1), Inches(5), Inches(4), data)
    path = tmp_path / "chart.pptx"
    prs.save(str(path))

    text = extract_content(str(path)).sections[0].text
    assert "【图表】" in text
    assert "| 类别 | 销售额 |" in text
    assert "| 第一季度 | 0.48 |" in text


def test_pptx_scatter_chart_without_categories(tmp_path):
    """XY/散点图没有类别轴，退化为「系列名: 数值」行而不是空表"""
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = XyChartData()
    series = data.add_series("性能点位")
    series.add_data_point(1.1, 2.2)
    series.add_data_point(1.5, 2.8)
    slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER,
                           Inches(1), Inches(1), Inches(5), Inches(4), data)
    path = tmp_path / "scatter.pptx"
    prs.save(str(path))

    text = extract_content(str(path)).sections[0].text
    assert "| 类别 |" not in text
    assert "性能点位: 2.2, 2.8" in text


def test_pptx_chart_block_not_used_as_title(tmp_path):
    """表格/图表块不能当章节标题，仍退化为页码"""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Q1"]
    data.add_series("S", (1,))
    slide.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(5), Inches(4), data)
    path = tmp_path / "chart_only.pptx"
    prs.save(str(path))

    assert extract_content(str(path)).sections[0].title == "第 1 页"


DIAGRAM_SAMPLE = r"E:\产品资料\产品资料\AnyBackup\AnyBackup Family 7 安可市场产品规划V2.pptx"


@pytest.mark.integration
@pytest.mark.skipif(not os.path.isfile(DIAGRAM_SAMPLE), reason="本地含 SmartArt 的样例胶片不存在")
def test_pptx_smartart_text():
    """SmartArt 图示的文字在 /ppt/diagrams/dataN.xml 里，不在形状文本框中"""
    ex = extract_content(DIAGRAM_SAMPLE)
    blocks = [s.text for s in ex.sections if "【图示】" in s.text]
    assert blocks, "图示文本未提取"
    assert any("AnyBackup Family 7" in text for text in blocks)


def test_formats_endpoint_exposes_registry():
    """GET /api/sources/formats 是注册表的对外出口，前端与 demo 脚本据此筛选文件"""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from api.routes.source import router

    app = FastAPI()
    app.include_router(router, prefix="/api")
    resp = TestClient(app).get("/api/sources/formats")

    assert resp.status_code == 200
    assert resp.json()["formats"] == list(supported_extensions())


# ── PDF（需要样例文件，找不到则跳过）──────────────────────

PDF_SAMPLE = r"E:\产品资料\技术方案文档\CDM技术资料\ESSS拓展分布式架构相关资料\ASU与AnyBackup CDM整合技术可行性探讨资料\Cohesity-SpanFS-and-SnapTree-WP.pdf"


@pytest.mark.integration
@pytest.mark.skipif(not os.path.isfile(PDF_SAMPLE), reason="本地样例 PDF 不存在")
def test_pdf_page_anchor():
    ex = extract_content(PDF_SAMPLE)
    assert ex.format == "pdf"
    assert ex.page_count and ex.page_count > 0
    assert ex.sections
    for root in ex.sections:
        assert root.page_start is not None
