"""
Word (.docx) 文档 PageIndex 树结构生成。

使用 python-docx 提取文本与标题样式，按 Word 内置 Heading 1-6 构建层级树，
复用 page_index_md 的树构建和摘要生成管线。
"""

import os
from docx import Document
from .utils import *
from .page_index_md import build_tree_from_nodes, generate_summaries_for_structure_md
from system.logger import logger


def extract_nodes_from_docx(docx_path: str):
    """从 Word 文档中提取标题节点和文本内容。

    通过段落样式识别标题（Heading 1-6），非标题段落作为前一个节点的正文。
    返回 (node_list, para_count, fallback_text)：
      - node_list: [{'title': str, 'level': int, 'text': str, 'para_index': int}, ...]
      - para_count: 文档总段落数
      - fallback_text: 全文纯文本（无标题时回退用）
    """
    doc = Document(docx_path)
    para_count = 0
    node_list = []
    all_texts = []

    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        para_count += 1
        style_name = (para.style.name if para.style else "")

        if text:
            all_texts.append(text)

        # 判断是否为 Word 内置标题样式
        is_heading = False
        level = 1
        if style_name.startswith("Heading") or style_name.startswith("heading"):
            try:
                level = int(style_name.replace("Heading", "").replace("heading", "").strip())
            except ValueError:
                level = 1
            is_heading = True
        elif style_name in ("Title", "title"):
            is_heading = True
            level = 1

        if is_heading and text:
            node_list.append({
                "title": text,
                "level": min(level, 6),
                "text": "",
                "para_index": idx,
            })

    # 为每个标题节点填充正文（两个标题之间的段落）
    for i, node in enumerate(node_list):
        start_idx = node["para_index"] + 1
        end_idx = node_list[i + 1]["para_index"] if i + 1 < len(node_list) else para_count

        body_lines = []
        for j in range(start_idx, end_idx):
            para_text = doc.paragraphs[j].text.strip()
            if para_text:
                body_lines.append(para_text)

        node["text"] = "\n\n".join(body_lines)

    fallback_text = "\n\n".join(all_texts)
    return node_list, para_count, fallback_text


async def docx_to_tree(docx_path: str,
                       if_update_title: bool = False,
                       if_add_node_summary: str = 'no',
                       summary_token_threshold: int = None,
                       model: str = None,
                       if_add_doc_description: str = 'no',
                       if_add_node_text: str = 'no',
                       if_add_node_id: str = 'yes'):
    """为 Word 文档生成 PageIndex 树结构。

    流程：python-docx 提取标题/正文 → 构建树 → 可选摘要/文档描述。

    Args:
        docx_path: .docx 文件路径
        if_update_title: 是否用 LLM 优化标题（实验性）
        if_add_node_summary: 是否生成节点摘要 ('yes'/'no')
        summary_token_threshold: 超过此 token 数才生成摘要
        model: LLM 模型名称
        if_add_doc_description: 是否生成文档描述
        if_add_node_text: 是否在输出中包含节点文本
        if_add_node_id: 是否生成节点 ID

    Returns:
        {'doc_name': str, 'para_count': int, 'structure': [...], 'doc_description': str}
    """
    logger.info(f"[DOCX] 提取节点: {docx_path}")

    node_list, para_count, fallback_text = extract_nodes_from_docx(docx_path)

    if not node_list:
        # 没有标题样式 → 回退：整个文档作为一个节点
        node_list.append({
            "title": os.path.splitext(os.path.basename(docx_path))[0],
            "text": fallback_text,
            "level": 1,
            "para_index": 0,
        })

    logger.info(f"[DOCX] 提取到 {len(node_list)} 个标题节点，共 {para_count} 个段落")

    # 可选：LLM 优化标题
    if if_update_title:
        logger.info("[DOCX] LLM 优化标题...")
        for node in node_list:
            node["title"] = await update_node_title(node, model=model)

    # 构建树结构
    logger.info("[DOCX] 构建树结构...")
    tree_structure = build_tree_from_nodes(node_list)

    if if_add_node_id == 'yes':
        write_node_id(tree_structure)

    # 生成摘要
    if if_add_node_summary == 'yes':
        tree_structure = format_structure(tree_structure,
                                          order=['title', 'node_id', 'para_index', 'summary', 'text', 'nodes'])

        logger.info("[DOCX] 生成节点摘要...")
        tree_structure = await generate_summaries_for_structure_md(
            tree_structure, summary_token_threshold=summary_token_threshold, model=model)

        if if_add_node_text == 'no':
            tree_structure = format_structure(tree_structure,
                                              order=['title', 'node_id', 'para_index', 'summary', 'nodes'])

        if if_add_doc_description == 'yes':
            logger.info("[DOCX] 生成文档描述...")
            clean_structure = create_clean_structure_for_description(tree_structure)
            doc_description = generate_doc_description(clean_structure, model=model)
            return {
                "doc_name": os.path.splitext(os.path.basename(docx_path))[0],
                "doc_description": doc_description,
                "para_count": para_count,
                "structure": tree_structure,
            }
    else:
        if if_add_node_text == 'yes':
            tree_structure = format_structure(tree_structure,
                                              order=['title', 'node_id', 'para_index', 'text', 'nodes'])
        else:
            tree_structure = format_structure(tree_structure,
                                              order=['title', 'node_id', 'para_index', 'nodes'])

    return {
        "doc_name": os.path.splitext(os.path.basename(docx_path))[0],
        "para_count": para_count,
        "structure": tree_structure,
    }