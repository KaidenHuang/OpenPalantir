"""PPTX 转换器 —— python-pptx：一页幻灯片一个章节，幻灯片序号作为页码锚点

章节标题依次取标题占位符（Title / Center Title）、页面第一行文本、"第 N 页"——
不少企业模板把标题画成普通文本框，只认占位符会让整套胶片的章节标题全部退化。
其余形状（正文占位符、副标题、文本框、自选图形、组合图形内的文本、表格、图表、
SmartArt 图示）按幻灯片内顺序进正文；演讲者备注并入同一章节正文而非拆子章节——
父节点在摘要树里只留 prefix_summary，自身正文不落盘，拆子章节会让幻灯片正文从
抽取批次里消失。

图表走 python-pptx 的 chart 对象；图示（SmartArt）的文字不在形状里而在
/ppt/diagrams/dataN.xml，需按 graphicFrame 的 r:dm 关系取该 part 的 dgm:pt 文本。
"""
import os
from typing import Iterator, List

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError

from system.logger import logger
from ..models import ExtractedDocument, Section

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
NOTES_PREFIX = "（备注）"
TABLE_PREFIX = "|"
GRAPHIC_PREFIX = "【"

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _iter_shapes(shapes) -> Iterator:
    """展开组合图形，按幻灯片内的物理顺序产出所有形状"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _is_title(shape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in _TITLE_TYPES


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _cell(value) -> str:
    """图表数值/类别原样转文本，不推断单位与百分号"""
    return "" if value is None else str(value).strip()


def _label(value, fallback: str) -> str:
    """系列/类别名可能为空，统一兜底成可读标签"""
    return (value or "").strip() or fallback


def _chart_text(chart) -> str:
    """图表 → 标题 + 类别×系列表格；XY/散点图没有类别轴（categories 为空），退化为系列数值行"""
    lines = [f"【图表】{_label(chart.chart_title.text_frame.text, '')}" if chart.has_title else "【图表】"]
    for plot in chart.plots:
        categories = [_cell(c) for c in plot.categories]
        series = list(plot.series)
        if categories:
            header = ["类别"] + [_label(s.name, f"系列{i + 1}") for i, s in enumerate(series)]
            lines.append("| " + " | ".join(header) + " |")
            for position, category in enumerate(categories):
                row = [category]
                for one in series:
                    values = list(one.values)
                    row.append(_cell(values[position]) if position < len(values) else "")
                lines.append("| " + " | ".join(row) + " |")
        else:
            for one in series:
                values = ", ".join(_cell(v) for v in one.values)
                lines.append(f"{_label(one.name, '系列')}: {values}")
    return "\n".join(lines)


def _diagram_text(shape, slide) -> str:
    """SmartArt 图示 → dgm:pt 里的文字；python-pptx 没有图示 API，需按 r:dm 关系读 diagramData part"""
    graphic_data = shape._element.find(f".//{{{_NS_A}}}graphicData")
    if graphic_data is None or graphic_data.get("uri") != _NS_DGM:
        return ""
    rel_ids = graphic_data.find(f"{{{_NS_DGM}}}relIds")
    data_rid = rel_ids.get(f"{{{_NS_R}}}dm") if rel_ids is not None else None
    if not data_rid:
        return ""
    root = etree.fromstring(slide.part.related_part(data_rid).blob)
    lines = []
    for point in root.iter(f"{{{_NS_DGM}}}pt"):
        texts = [node.text.strip() for node in point.iter(f"{{{_NS_A}}}t") if node.text and node.text.strip()]
        if texts:
            lines.append(" ".join(texts))
    return "\n".join(["【图示】"] + lines) if lines else ""


def _text_frame_lines(text_frame) -> List[str]:
    """段落 → 文本行；大纲层级 0 保持整句，更深层级渲染成缩进列表以保留从属关系"""
    lines = []
    for paragraph in text_frame.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = paragraph.level or 0
        lines.append(text if level == 0 else f"{'  ' * level}- {text}")
    return lines


def _placeholder_title(slide) -> str:
    for shape in slide.shapes:
        if _is_title(shape):
            return shape.text_frame.text.strip() if shape.has_text_frame else ""
    return ""


def _body_blocks(slide) -> List[str]:
    """幻灯片正文文本块：表格/图表/图示整块产出，文本逐段产出（标题占位符已排除）"""
    blocks: List[str] = []
    for shape in _iter_shapes(slide.shapes):
        if _is_title(shape):
            continue
        if shape.has_table:
            blocks.append(_table_text(shape.table))
        elif shape.has_chart:
            blocks.append(_chart_text(shape.chart))
        else:
            diagram = _diagram_text(shape, slide)
            if diagram:
                blocks.append(diagram)
            elif shape.has_text_frame:
                blocks.extend(_text_frame_lines(shape.text_frame))
    return [block for block in blocks if block and not block.isspace()]


def _notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    text_frame = slide.notes_slide.notes_text_frame
    if text_frame is None:
        return ""
    return "\n".join(_text_frame_lines(text_frame))


def _take_text_title(blocks: List[str]) -> str:
    """无标题占位符时取第一个真实文本行作标题，并从正文中摘除（表格/图表/图示块不做标题）"""
    for position, block in enumerate(blocks):
        if block.startswith(TABLE_PREFIX) or block.startswith(GRAPHIC_PREFIX):
            continue
        return blocks.pop(position)
    return ""


def extract_pptx(path: str) -> ExtractedDocument:
    try:
        presentation = Presentation(path)
    except PackageNotFoundError as e:
        raise ValueError(f"PPTX 文件无法读取（旧版 .ppt 二进制格式不支持，请先另存为 .pptx）: {e}")

    doc_name = os.path.splitext(os.path.basename(path))[0]
    sections: List[Section] = []
    skipped = 0

    for index, slide in enumerate(presentation.slides, start=1):
        blocks = _body_blocks(slide)
        title = _placeholder_title(slide) or _take_text_title(blocks)

        notes = _notes_text(slide)
        if notes:
            blocks.append(f"{NOTES_PREFIX}{notes}")
        text = "\n".join(blocks)

        if not title and not text:
            skipped += 1
            continue
        # 空正文的叶子节点不会进摘要树与抽取批次，故纯标题页把标题回填进正文
        sections.append(Section(title=title or f"第 {index} 页", level=1,
                                text=text or title,
                                page_start=index, page_end=index))

    page_count = len(presentation.slides)
    if skipped:
        logger.info(f"[pptx_converter] {skipped} 页无可提取文本（纯图片页），已跳过: {doc_name}")
    if not sections:
        raise ValueError("PPTX 无可提取的文本内容")

    return ExtractedDocument(doc_name=doc_name, format="pptx",
                             page_count=page_count, sections=sections)
